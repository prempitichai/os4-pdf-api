#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# VERSION: docx-template-v1
"""
generate_bg_withdraw.py
=======================
Railway API endpoints — ใช้ Word template + python-docx fill + LibreOffice convert
  1. /generate_bg_withdraw  — หนังสือแจ้งขอถอนหลักประกันสัญญา
  2. /generate_bg_poa       — หนังสือมอบอำนาจขอคืนหนังสือค้ำประกัน

Templates:
  bg_withdraw_template.docx  — จดหมายขอถอนหลักประกัน.docx (ใส่ placeholder แล้ว)
  bg_poa_template.docx       — จดหมายมอบอำนาจ.docx (ใส่ placeholder แล้ว)

Placeholders ในไฟล์ docx:
  {contractId}, {contractName}, {signedDate}, {company}
  {guaranteeNumber}, {guaranteeIssueDate}, {guaranteeValue}
  {signerName}, {signerPosition}, {docDate}, {docNumber}
  {grantorName}, {grantorIdCard}, {granteeName}, {granteeIdCard}
  {granteeAddress}, {witness1}, {witness2}
"""

import os, base64, logging, subprocess, tempfile, shutil
from flask import request, jsonify

logger = logging.getLogger(__name__)

# ── Template paths ─────────────────────────────────────────────────────
def _tpl(name):
    base = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base, name)

# ── Number formatter ───────────────────────────────────────────────────
def _fmt(val):
    if not val: return ''
    try:
        n = float(str(val).replace(',', ''))
        return f'{n:,.2f}'
    except Exception:
        return str(val)

# ── Fill placeholders in docx ──────────────────────────────────────────
def _fill_docx(template_path, replacements):
    """
    โหลด Word template → replace {placeholder} ด้วยค่าจริง → คืน bytes
    รองรับ placeholder ที่ถูกตัดข้าม run ใน XML ด้วยการ merge runs ก่อน
    """
    from docx import Document
    from docx.oxml.ns import qn
    from copy import deepcopy
    import re

    doc = Document(template_path)

    def _replace_in_para(para):
        """Replace placeholder ใน paragraph — จัดการ run ที่ถูกตัดได้"""
        # รวม text ทั้งหมดใน paragraph
        full_text = ''.join(r.text for r in para.runs)
        # ตรวจว่ามี placeholder ไหม
        has_placeholder = any(k in full_text for k in replacements)
        if not has_placeholder:
            return
        # แทนที่ทั้งหมด
        new_text = full_text
        for k, v in replacements.items():
            new_text = new_text.replace(k, str(v or ''))
        # ใส่ text ใหม่เข้า run แรก ลบ run ที่เหลือ
        if para.runs:
            para.runs[0].text = new_text
            for r in para.runs[1:]:
                r.text = ''

    # แทนที่ใน body
    for para in doc.paragraphs:
        _replace_in_para(para)

    # แทนที่ใน table
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    _replace_in_para(para)

    # แทนที่ใน header/footer
    for section in doc.sections:
        for para in section.header.paragraphs:
            _replace_in_para(para)
        for para in section.footer.paragraphs:
            _replace_in_para(para)

    # บันทึกเป็น bytes
    import io
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── Convert docx → PDF ────────────────────────────────────────────────
def _docx_to_pdf(docx_bytes, filename='output.docx'):
    """
    แปลง docx bytes → PDF bytes ผ่าน LibreOffice
    """
    with tempfile.TemporaryDirectory() as tmpdir:
        # เขียน docx ลงไฟล์ชั่วคราว
        docx_path = os.path.join(tmpdir, filename)
        with open(docx_path, 'wb') as f:
            f.write(docx_bytes)

        # รัน LibreOffice
        result = subprocess.run(
            ['libreoffice', '--headless', '--convert-to', 'pdf',
             '--outdir', tmpdir, docx_path],
            capture_output=True, text=True, timeout=60
        )

        if result.returncode != 0:
            raise RuntimeError(f'LibreOffice error: {result.stderr}')

        # หา PDF ที่ได้
        pdf_name = filename.replace('.docx', '.pdf')
        pdf_path = os.path.join(tmpdir, pdf_name)
        if not os.path.exists(pdf_path):
            # ลอง find
            pdfs = [f for f in os.listdir(tmpdir) if f.endswith('.pdf')]
            if not pdfs:
                raise RuntimeError('PDF not found after LibreOffice conversion')
            pdf_path = os.path.join(tmpdir, pdfs[0])

        with open(pdf_path, 'rb') as f:
            return f.read()


# ══════════════════════════════════════════════════════════════════════
# ENDPOINT 1: หนังสือแจ้งขอถอนหลักประกันสัญญา
# ══════════════════════════════════════════════════════════════════════
def generate_bg_withdraw():
    try:
        data = request.get_json(force=True) or {}

        contract_id   = str(data.get('contractId', '') or '')
        contract_name = str(data.get('contractName', '') or '')
        signed_date   = str(data.get('signedDate', '') or '')
        company       = str(data.get('company', '') or '')
        bg_number     = str(data.get('guaranteeNumber', '') or '')
        bg_date       = str(data.get('guaranteeIssueDate', '') or '')
        bg_value      = _fmt(data.get('guaranteeValue', ''))
        signer        = str(data.get('signerName', '') or '')
        signer_pos    = str(data.get('signerPosition', 'Corporate Lawyers') or 'Corporate Lawyers')
        doc_number    = str(data.get('docNumber', '') or contract_id)
        doc_date      = str(data.get('docDate', '') or '')
        entity        = data.get('entity') or {}
        entity_name   = entity.get('name', 'บริษัท เอส ซี เอ็ม เทคโนโลจีส์ จำกัด')

        replacements = {
            '{contractId}':          contract_id,
            '{contractName}':        contract_name,
            '{signedDate}':          signed_date,
            '{company}':             company,
            '{guaranteeNumber}':     bg_number,
            '{guaranteeIssueDate}':  bg_date,
            '{guaranteeValue}':      bg_value,
            '{signerName}':          signer,
            '{signerPosition}':      signer_pos,
            '{docNumber}':           doc_number,
            '{docDate}':             doc_date,
            '{entityName}':          entity_name,
            # backward compat
            '(เลขที่สัญญา)':        doc_number,
        }

        tpl_path  = _tpl('bg_withdraw_template.docx')
        docx_bytes = _fill_docx(tpl_path, replacements)
        pdf_bytes  = _docx_to_pdf(docx_bytes, 'withdraw.docx')
        file_name  = f'หนังสือขอถอนหลักประกัน_{contract_id.replace("/","-")}.pdf'

        return jsonify({
            'success':   True,
            'pdfBase64': base64.b64encode(pdf_bytes).decode(),
            'fileName':  file_name
        })
    except Exception as e:
        logger.error(f'bg_withdraw error: {e}')
        return jsonify({'success': False, 'message': str(e)}), 500


# ══════════════════════════════════════════════════════════════════════
# ENDPOINT 2: หนังสือมอบอำนาจ
# ══════════════════════════════════════════════════════════════════════
def generate_bg_poa():
    try:
        data = request.get_json(force=True) or {}

        contract_id  = str(data.get('contractId', '') or '')
        bg_number    = str(data.get('guaranteeNumber', '') or '')
        bg_value     = _fmt(data.get('guaranteeValue', ''))
        company      = str(data.get('company', '') or '')
        doc_date     = str(data.get('docDate', '') or '')
        entity       = data.get('entity') or {}
        grantor      = data.get('grantor') or {}
        grantee      = data.get('grantee') or {}
        witnesses    = data.get('witnesses') or [{}, {}]
        while len(witnesses) < 2: witnesses.append({})

        entity_name  = entity.get('name', 'บริษัท เอส ซี เอ็ม เทคโนโลจีส์ จำกัด')
        grantor_name = str(grantor.get('name', 'นายบัณฑิต หมั้นทรัพย์') or '')
        grantor_id   = str(grantor.get('idCard', '3101201187013') or '')
        grantee_name = str(grantee.get('name', '') or '')
        grantee_id   = str(grantee.get('idCard', '') or '')
        grantee_addr = str(grantee.get('address', '') or '')
        w1           = str(witnesses[0].get('name', 'นายปิติชัย พัฒนกิจกุล') or '')
        w2           = str(witnesses[1].get('name', 'นางสาวสิริกาญจนา จันทร์อ่อน') or '')

        replacements = {
            '{contractId}':      contract_id,
            '{guaranteeNumber}': bg_number,
            '{guaranteeValue}':  bg_value,
            '{company}':         company,
            '{docDate}':         doc_date,
            '{entityName}':      entity_name,
            '{grantorName}':     grantor_name,
            '{grantorIdCard}':   grantor_id,
            '{granteeName}':     grantee_name,
            '{granteeIdCard}':   grantee_id,
            '{granteeAddress}':  grantee_addr,
            '{witness1}':        w1,
            '{witness2}':        w2,
            # ข้อความที่ fix อยู่ใน template แล้ว แต่ใส่ไว้เผื่อ
            '(เลขที่สัญญา)':    f'({contract_id})',
            'วันที่':            f'วันที่ {doc_date}',
            # ชื่อที่ hardcode ใน template — replace ด้วย dynamic
            'นายบัณฑิต หมั้นทรัพย์':           grantor_name,
            'นายวุฒิชัย บันลือทรัพย์':          grantee_name or 'นายวุฒิชัย บันลือทรัพย์',
            'บัตรประชาชนเลขที่ 3101201187013':   f'บัตรประชาชนเลขที่ {grantor_id}',
            'ผู้ถือบัตรประชาชนเลขที่ 3320101225506': f'ผู้ถือบัตรประชาชนเลขที่ {grantee_id}',
            'นายปิติชัย พัฒนกิจกุล':            w1,
            'นางสาวสิริกาญจนา จันทร์อ่อน':      w2,
        }

        tpl_path   = _tpl('bg_poa_template.docx')
        docx_bytes = _fill_docx(tpl_path, replacements)
        pdf_bytes  = _docx_to_pdf(docx_bytes, 'poa.docx')
        file_name  = f'หนังสือมอบอำนาจ_{contract_id.replace("/","-")}.pdf'

        return jsonify({
            'success':   True,
            'pdfBase64': base64.b64encode(pdf_bytes).decode(),
            'fileName':  file_name
        })
    except Exception as e:
        logger.error(f'bg_poa error: {e}')
        return jsonify({'success': False, 'message': str(e)}), 500


# ══════════════════════════════════════════════════════════════════════
# Registration
# ══════════════════════════════════════════════════════════════════════
def register_bg_withdraw_routes(app):
    app.add_url_rule('/generate_bg_withdraw', 'generate_bg_withdraw',
                     generate_bg_withdraw, methods=['POST'])
    app.add_url_rule('/generate_bg_poa', 'generate_bg_poa',
                     generate_bg_poa, methods=['POST'])
